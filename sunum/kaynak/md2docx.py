#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown -> profesyonel .docx dönüştürücü (TEKNOFEST Büke jüri dökümanları).

Desteklenen markdown alt kümesi:
    # H1 (belge başlığı, kapakta kullanılır)   ## H2   ### H3   #### H4
    paragraf, **kalın**, *italik*, `kod`
    - madde  /  1. numaralı madde
    | pipe | tablo |
    > alıntı kutusu
    ```kod bloğu```
    --- yatay çizgi
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------- renk paleti
BRAND = RGBColor(0xE1, 0x4A, 0x1E)        # Hepsiburada turuncusu
INK = RGBColor(0x1A, 0x1D, 0x23)
MUTED = RGBColor(0x5A, 0x63, 0x72)
RULE = "D8DCE3"
CODE_BG = "F4F5F7"
QUOTE_BG = "FFF4EE"
HEAD_BG = "2B303B"

BODY_FONT = "Calibri"
HEAD_FONT = "Calibri Light"
MONO_FONT = "Consolas"


# ---------------------------------------------------------------- xml yardımcıları
def _shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _para_shading(paragraph, fill: str) -> None:
    _shade(paragraph._p.get_or_add_pPr(), fill)


def _cell_shading(cell, fill: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), fill)


def _para_border(paragraph, *, edge: str, size: int = 6, color: str = RULE,
                 space: int = 4) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    borders.append(el)


def _keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    p_pr.append(el)


def _no_widow_table_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    tr_pr.append(el)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)

    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)

    if placeholder:
        paragraph.add_run(placeholder)

    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _update_fields_on_open(document) -> None:
    """İçindekiler tablosunu Word açılışta güncellesin.

    CT_Settings şemasında ``w:updateFields``, ``w:compat``ten **önce** gelir;
    sona eklemek Word'ün "okunamayan içerik" uyarısına yol açabilir.
    """
    settings = document.settings.element
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    for anchor in ("w:compat", "w:rsids", "w:themeFontLang"):
        node = settings.find(qn(anchor))
        if node is not None:
            node.addprevious(el)
            return
    settings.append(el)


# ---------------------------------------------------------------- stil kurulumu
def build_styles(document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.space_after = Pt(7)
    pf.line_spacing = 1.18

    specs = [
        ("Heading 1", 19, True, BRAND, 20, 8),
        ("Heading 2", 14.5, True, INK, 15, 5),
        ("Heading 3", 12, True, MUTED, 12, 4),
        ("Heading 4", 10.5, True, MUTED, 10, 3),
    ]
    for name, size, bold, color, before, after in specs:
        st = document.styles[name]
        st.font.name = HEAD_FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = color
        st.font.italic = False
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        st = document.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(10.5)
        st.font.color.rgb = INK
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.15


# ---------------------------------------------------------------- satır içi biçim
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`)")


def add_inline(paragraph, text: str, *, base_bold=False, base_italic=False,
               color=None, size=None, mono=False):
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        bold, italic, code = base_bold, base_italic, mono
        body = chunk
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            bold, body = True, chunk[2:-2]
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            italic, body = True, chunk[1:-1]
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            code, body = True, chunk[1:-1]
        run = paragraph.add_run(body)
        run.bold = bold
        run.italic = italic
        run.font.name = MONO_FONT if code else (BODY_FONT if not mono else MONO_FONT)
        if code:
            run.font.size = Pt((size or 10.5) - 1)
            run.font.color.rgb = RGBColor(0xB0, 0x33, 0x0E)
        else:
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
    return paragraph


# ---------------------------------------------------------------- blok yazıcılar
def write_table(document, rows):
    header, body = rows[0], rows[1:]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    hdr = table.rows[0]
    _repeat_header(hdr)
    _no_widow_table_row(hdr)
    for cell, text in zip(hdr.cells, header):
        _cell_shading(cell, HEAD_BG)
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(3)
        para.paragraph_format.space_after = Pt(3)
        add_inline(para, text, base_bold=True,
                   color=RGBColor(0xFF, 0xFF, 0xFF), size=9.5)

    for index, line in enumerate(body):
        row = table.add_row()
        _no_widow_table_row(row)
        if index % 2 == 1:
            for cell in row.cells:
                _cell_shading(cell, "F7F8FA")
        for cell, text in zip(row.cells, line):
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, text, size=9.5)

    tail = document.add_paragraph()
    tail.paragraph_format.space_after = Pt(6)
    tail.paragraph_format.space_before = Pt(0)
    return table


def write_code(document, lines):
    for index, line in enumerate(lines):
        para = document.add_paragraph()
        pf = para.paragraph_format
        pf.space_before = Pt(6 if index == 0 else 0)
        pf.space_after = Pt(6 if index == len(lines) - 1 else 0)
        pf.left_indent = Cm(0.4)
        pf.line_spacing = 1.0
        _para_shading(para, CODE_BG)
        run = para.add_run(line if line.strip() else " ")
        run.font.name = MONO_FONT
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x24, 0x29, 0x33)


def write_quote(document, text):
    para = document.add_paragraph()
    pf = para.paragraph_format
    pf.left_indent = Cm(0.5)
    pf.right_indent = Cm(0.3)
    pf.space_before = Pt(6)
    pf.space_after = Pt(8)
    _para_shading(para, QUOTE_BG)
    _para_border(para, edge="left", size=18, color="E14A1E", space=6)
    add_inline(para, text, size=10)


def write_rule(document):
    para = document.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(8)
    _para_border(para, edge="bottom", size=6, color=RULE, space=1)


# ---------------------------------------------------------------- kapak + TOC
def cover_page(document, title, subtitle, meta_rows):
    for _ in range(3):
        document.add_paragraph()

    tag = document.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = tag.add_run("TEKNOFEST 2026  ·  HEPSİBURADA")
    run.font.name = HEAD_FONT
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = BRAND

    head = document.add_paragraph()
    head.paragraph_format.space_before = Pt(6)
    head.paragraph_format.space_after = Pt(4)
    run = head.add_run(title)
    run.font.name = HEAD_FONT
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = INK

    sub = document.add_paragraph()
    sub.paragraph_format.space_after = Pt(16)
    run = sub.add_run(subtitle)
    run.font.name = HEAD_FONT
    run.font.size = Pt(14)
    run.font.color.rgb = MUTED

    write_rule(document)

    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in meta_rows:
        row = table.add_row()
        _cell_shading(row.cells[0], "F7F8FA")
        left = row.cells[0].paragraphs[0]
        add_inline(left, key, base_bold=True, size=10, color=MUTED)
        right = row.cells[1].paragraphs[0]
        add_inline(right, value, size=10)
    for row in table.rows:
        row.cells[0].width = Cm(5.4)
        row.cells[1].width = Cm(11.0)

    document.add_paragraph()


def toc_page(document):
    head = document.add_paragraph()
    head.paragraph_format.space_after = Pt(10)
    run = head.add_run("İçindekiler")
    run.font.name = HEAD_FONT
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = BRAND

    para = document.add_paragraph()
    _field(para, r' TOC \o "1-3" \h \z \u ',
           "Word'de içindekiler tablosunu güncellemek için: tabloya sağ tıklayın "
           "→ “Alanı Güncelleştir” (veya Ctrl+A, ardından F9).")


def page_setup(document, doc_title):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.0)
    return section


def decorate_section(section, doc_title):
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run(f"{doc_title}   ·   Takım Büke")
    run.font.name = HEAD_FONT
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    _para_border(header, edge="bottom", size=4, color=RULE, space=2)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Sayfa ")
    run.font.name = HEAD_FONT
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    _field(footer, " PAGE ", "1")
    run = footer.add_run(" / ")
    run.font.name = HEAD_FONT
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    _field(footer, " NUMPAGES ", "1")
    for r in footer.runs:
        r.font.name = HEAD_FONT
        r.font.size = Pt(8.5)
        r.font.color.rgb = MUTED


# ---------------------------------------------------------------- ayrıştırıcı
def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_divider(line):
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells)


def convert(md_path: Path, docx_path: Path, subtitle: str, meta_rows):
    raw = md_path.read_text(encoding="utf-8").replace("\r\n", "\n")
    lines = raw.split("\n")

    title = md_path.stem
    for line in lines:
        if line.startswith("# "):
            title = line[2:].strip()
            break

    document = Document()
    build_styles(document)
    section = page_setup(document, title)

    cover_page(document, title, subtitle, meta_rows)
    document.add_page_break()
    toc_page(document)
    document.add_page_break()
    decorate_section(section, title)

    index = 0
    total = len(lines)
    seen_h1 = False
    while index < total:
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        # kod bloğu
        if stripped.startswith("```"):
            index += 1
            block = []
            while index < total and not lines[index].strip().startswith("```"):
                block.append(lines[index].rstrip())
                index += 1
            index += 1
            while block and not block[0].strip():
                block.pop(0)
            while block and not block[-1].strip():
                block.pop()
            if block:
                write_code(document, block)
            continue

        # tablo
        if stripped.startswith("|") and index + 1 < total and is_divider(lines[index + 1]):
            header = split_table_row(lines[index])
            index += 2
            body = []
            while index < total and lines[index].strip().startswith("|"):
                cells = split_table_row(lines[index])
                cells = (cells + [""] * len(header))[:len(header)]
                body.append(cells)
                index += 1
            write_table(document, [header] + body)
            continue

        # yatay çizgi
        if re.fullmatch(r"-{3,}|_{3,}|\*{3,}", stripped):
            write_rule(document)
            index += 1
            continue

        # başlıklar
        match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if match:
            level, text = len(match.group(1)), match.group(2).strip()
            if level == 1:
                if seen_h1:
                    para = document.add_paragraph(style="Heading 1")
                    add_inline(para, text)
                seen_h1 = True
                index += 1
                continue
            style = f"Heading {min(level - 1, 4)}"
            para = document.add_paragraph(style=style)
            add_inline(para, text)
            _keep_with_next(para)
            index += 1
            continue

        # alıntı
        if stripped.startswith(">"):
            block = []
            while index < total and lines[index].strip().startswith(">"):
                block.append(lines[index].strip().lstrip(">").strip())
                index += 1
            write_quote(document, " ".join(p for p in block if p))
            continue

        # madde listesi
        bullet = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if bullet:
            indent = len(bullet.group(1)) // 2
            para = document.add_paragraph(style="List Bullet")
            if indent:
                para.paragraph_format.left_indent = Cm(0.75 + 0.6 * indent)
            add_inline(para, bullet.group(2).strip())
            index += 1
            continue

        number = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if number:
            indent = len(number.group(1)) // 2
            para = document.add_paragraph(style="List Number")
            if indent:
                para.paragraph_format.left_indent = Cm(0.75 + 0.6 * indent)
            add_inline(para, number.group(2).strip())
            index += 1
            continue

        # paragraf (ardışık satırları birleştir)
        block = [stripped]
        index += 1
        while index < total:
            nxt = lines[index]
            s = nxt.strip()
            if (not s or s.startswith(("#", ">", "|", "```"))
                    or re.match(r"^\s*[-*+]\s+", nxt)
                    or re.match(r"^\s*\d+[.)]\s+", nxt)
                    or re.fullmatch(r"-{3,}|_{3,}|\*{3,}", s)):
                break
            block.append(s)
            index += 1
        para = document.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(para, " ".join(block))

    core = document.core_properties
    core.title = title
    core.subject = subtitle
    core.author = "Takım Büke"
    core.category = "TEKNOFEST 2026 · Hepsiburada Lojistik Optimizasyonu"
    core.comments = "Final Backtest aşaması jüri dökümanı"

    _update_fields_on_open(document)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(docx_path))
    return title


if __name__ == "__main__":
    src, dst, sub = sys.argv[1], sys.argv[2], sys.argv[3]
    rows = [
        ("Yarışma", "TEKNOFEST 2026 · Hepsiburada"),
        ("Kategori", "Yapay Zekâ Destekli Lojistik Anahat Optimizasyonu"),
        ("Aşama", "Final — Backtest"),
        ("Takım", "Büke"),
        ("Belge", sub),
    ]
    name = convert(Path(src), Path(dst), sub, rows)
    print(f"OK  {name}  ->  {dst}")
