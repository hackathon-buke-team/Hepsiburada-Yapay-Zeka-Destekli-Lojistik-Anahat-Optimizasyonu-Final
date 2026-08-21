"""Doğrulanmış forecast ve plan DataFrame'lerini Excel'e aktarır."""
from __future__ import annotations

import os
import shutil
import tempfile
from datetime import time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import load_workbook

from src.schemas import (FORECAST_COLS, PLAN_COLS, _canonical_forecast_time,
                         validate_forecast_grid, validate_plan)


def _validation_error(label: str, errors: list) -> ValueError:
    details = "\n".join(f"- {error}" for error in errors)
    return ValueError(f"{label} doğrulaması başarısız:\n{details}")


def _format_forecast_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook.active
    slot_col = FORECAST_COLS.index("Talep Tamamlama Saati") + 1
    desi_col = FORECAST_COLS.index("Tahmin Edilen Desi") + 1

    for row in range(2, sheet.max_row + 1):
        slot_cell = sheet.cell(row=row, column=slot_col)
        canonical = _canonical_forecast_time(slot_cell.value)
        if canonical is None:
            raise ValueError(
                f"Geçersiz forecast saati: {slot_cell.value!r}")
        hour, minute = (int(part) for part in canonical.split(":"))
        slot_cell.value = time(hour, minute)
        slot_cell.number_format = "h:mm"
        sheet.cell(row=row, column=desi_col).number_format = "0.000"

    workbook.save(path)


def _write_round_trip(df: pd.DataFrame, path, columns: list,
                      validator, label: str,
                      formatter: Callable[[Path], None] | None = None) -> Path:
    errors = validator(df)
    if errors:
        raise _validation_error(label, errors)

    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}-", suffix=".xlsx",
        dir=output_path.parent)
    os.close(fd)
    temporary_path = Path(temporary_name)

    try:
        df.loc[:, columns].to_excel(temporary_path, index=False)
        if formatter is not None:
            formatter(temporary_path)
        round_tripped = pd.read_excel(temporary_path)
        errors = validator(round_tripped)
        if errors:
            raise _validation_error(f"{label} round-trip", errors)
        os.replace(temporary_path, output_path)
    finally:
        temporary_path.unlink(missing_ok=True)

    return output_path


def write_forecast_xlsx(df: pd.DataFrame, path, data, start, end,
                        ods=None) -> Path:
    """Forecast'u grid dahil önce ve Excel round-trip sonrasında doğrular."""
    return _write_round_trip(
        df, path, FORECAST_COLS,
        lambda frame: validate_forecast_grid(
            frame, data, start, end, ods=ods),
        "Forecast",
        formatter=_format_forecast_workbook)


def require_plan_total(df: pd.DataFrame, expected_total: float,
                       tolerance: float = 0.01) -> None:
    def finite_decimal(value, error_message: str) -> Decimal:
        try:
            number = Decimal(str(value))
        except (InvalidOperation, OverflowError, TypeError, ValueError) as error:
            raise ValueError(error_message) from error
        if not number.is_finite():
            raise ValueError(error_message)
        return number

    declared_error = (
        "Toplam maliyet değerlerinin her biri sonlu bir sayı olmalı")
    try:
        converted = pd.to_numeric(df["Toplam maliyet"], errors="raise")
    except (OverflowError, TypeError, ValueError) as error:
        raise ValueError(declared_error) from error

    declared = sum(
        (finite_decimal(value, declared_error) for value in converted),
        Decimal("0"))
    referee = finite_decimal(
        expected_total, "Hakem toplam maliyeti sonlu bir sayı olmalı")
    tolerance_error = (
        "Tolerans sonlu ve negatif olmayan bir sayı olmalı")
    allowed_difference = finite_decimal(tolerance, tolerance_error)
    if allowed_difference < 0:
        raise ValueError(tolerance_error)

    if abs(declared - referee) > allowed_difference:
        raise ValueError(
            "Toplam maliyet uzlaşmıyor: "
            f"beyan {declared:.2f} TL, hakem {referee:.2f} TL"
        )


def _format_plan_workbook(path: Path) -> None:
    """Match the official plan template's 0.000 desi cell format."""
    workbook = load_workbook(path)
    sheet = workbook.active
    desi_col = PLAN_COLS.index("Taşınan Desi") + 1

    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=desi_col).number_format = "0.000"

    workbook.save(path)


def write_plan_xlsx(df: pd.DataFrame, path, data) -> Path:
    """Planı şema doğrulamasından geçirerek template kolonlarıyla yazar."""
    return _write_round_trip(
        df, path, PLAN_COLS,
        lambda frame: validate_plan(frame, data),
        "Plan",
        formatter=_format_plan_workbook)


def publish_workbooks(
        staged_plan: str | Path,
        staged_forecast: str | Path,
        plan_destination: str | Path,
        forecast_destination: str | Path) -> tuple[Path, Path]:
    """Publish plan first and restore both destinations on ordinary failure."""
    sources = (Path(staged_plan), Path(staged_forecast))
    destinations = (Path(plan_destination), Path(forecast_destination))

    for source in sources:
        if source.is_symlink():
            raise ValueError(f"Staged workbook symlink olamaz: {source}")
    missing_sources = [str(source) for source in sources if not source.is_file()]
    if missing_sources:
        raise FileNotFoundError(
            "Staged workbook bulunamadı: " + ", ".join(missing_sources))
    for destination in destinations:
        if destination.is_symlink():
            raise ValueError(f"Official workbook symlink olamaz: {destination}")
        if destination.exists() and not destination.is_file():
            raise ValueError(
                f"Official workbook regular dosya olmalı: {destination}")

    resolved_paths = [path.resolve() for path in (*sources, *destinations)]
    if len(set(resolved_paths)) != 4:
        raise ValueError("Staged ve official workbook yolları farklı olmalı")
    if destinations[0].parent.resolve() != destinations[1].parent.resolve():
        raise ValueError("Official workbook hedefleri aynı dizinde olmalı")

    existing_paths = [
        path for path in (*sources, *destinations) if path.exists()
    ]
    for index, path in enumerate(existing_paths):
        for other in existing_paths[index + 1:]:
            if path.samefile(other):
                raise ValueError(
                    f"Workbook yolları hardlink alias olamaz: {path}, {other}")
    for path in existing_paths:
        if path.stat().st_nlink > 1:
            raise ValueError(f"Workbook multi-link regular dosya olamaz: {path}")

    official_parent = destinations[0].parent
    official_parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
            prefix=".publish-", dir=official_parent) as backup_name:
        backup_dir = Path(backup_name)
        existed = tuple(destination.exists() for destination in destinations)
        backups = tuple(
            backup_dir / f"{index}.backup"
            for index in range(len(destinations))
        )
        for destination, backup, was_present in zip(
                destinations, backups, existed):
            if was_present:
                shutil.copy2(destination, backup)

        try:
            os.replace(sources[0], destinations[0])
            os.replace(sources[1], destinations[1])
        except Exception as publication_error:
            rollback_errors = []
            for destination, backup, was_present in zip(
                    destinations, backups, existed):
                try:
                    if was_present:
                        os.replace(backup, destination)
                    else:
                        destination.unlink(missing_ok=True)
                except Exception as rollback_error:
                    rollback_errors.append(
                        f"{destination}: {rollback_error!r}")
            if rollback_errors:
                raise RuntimeError(
                    "Workbook rollback başarısız: "
                    + "; ".join(rollback_errors)
                ) from publication_error
            raise

    return destinations
