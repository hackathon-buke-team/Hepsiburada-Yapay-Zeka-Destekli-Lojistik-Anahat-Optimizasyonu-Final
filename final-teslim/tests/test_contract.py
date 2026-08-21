"""``src.contract`` — Bölüm 4/5/7 girdi-çıktı sözleşmesi testleri."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
import pytest

from src import contract
from src.contract import (INPUT_ENV_VAR, InputContractError, horizon_days,
                          read_demand_table, resolve_input_path,
                          restore_demand_ids, write_final_plan)
from src.data import load_all
from src.schemas import FORECAST_COLS, PLAN_COLS

REFERENCE_INPUT = Path(__file__).resolve().parents[1] / "data" / "one_week_backtest.xlsx"


@pytest.fixture(scope="module")
def data():
    return load_all(Path(__file__).resolve().parents[1] / "datas",
                    with_demand=False)


def _write(frame: pd.DataFrame, path: Path) -> Path:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Sheet1")
    return path


def _demand_frame(rows: list) -> pd.DataFrame:
    return pd.DataFrame(rows, columns=FORECAST_COLS)


# ---------------------------------------------------------------- saat/tarih

@pytest.mark.parametrize("value,expected", [
    (dt.time(9, 0), "09:00"),
    (dt.time(17, 0), "17:00"),
    ("09:00", "09:00"),
    ("09:00:00", "09:00"),
    ("17:00:00", "17:00"),
    (" 17:00 ", "17:00"),
    (dt.datetime(2026, 6, 29, 17, 0), "17:00"),
    (pd.Timestamp("2026-06-29 09:00"), "09:00"),
    (0.375, "09:00"),                      # Excel gün kesri
    (pd.Timedelta(hours=17), "17:00"),
])
def test_to_hhmm_kabul_edilen_bicimler(value, expected):
    assert contract._to_hhmm(value) == expected


def test_to_hhmm_saniye_artigini_yukari_yuvarlar():
    # Yarışmanın "en yakın büyük tam sayıya yuvarla" kuralıyla tutarlı.
    assert contract._to_hhmm("09:00:30") == "09:01"


def test_to_hhmm_cozumlenemeyen_deger_none():
    assert contract._to_hhmm("yok") is None
    assert contract._to_hhmm(None) is None


@pytest.mark.parametrize("value,expected", [
    ("29.06.2026", dt.date(2026, 6, 29)),
    ("2026-06-29", dt.date(2026, 6, 29)),
    ("29/06/2026", dt.date(2026, 6, 29)),
    (dt.date(2026, 6, 29), dt.date(2026, 6, 29)),
    (dt.datetime(2026, 6, 29, 9, 0), dt.date(2026, 6, 29)),
    (pd.Timestamp("2026-06-29"), dt.date(2026, 6, 29)),
    (46202, dt.date(2026, 6, 29)),         # Excel seri numarası
])
def test_to_date_kabul_edilen_bicimler(value, expected):
    assert contract._to_date(value) == expected


@pytest.mark.parametrize("value,expected", [
    (5, 5), (5.4, 5), (5.6, 6), ("7", 7), (-3, 0), (None, None), ("abc", None),
])
def test_to_desi(value, expected):
    assert contract._to_desi(value) == expected


# --------------------------------------------------------------- girdi yolu

def test_resolve_input_path_ortam_degiskenini_once_dener(tmp_path, monkeypatch):
    target = tmp_path / "girdi.xlsx"
    target.write_bytes(b"x")
    monkeypatch.setenv(INPUT_ENV_VAR, str(target))
    assert resolve_input_path(tmp_path) == target


def test_resolve_input_path_yedek_yola_duser(tmp_path, monkeypatch):
    monkeypatch.delenv(INPUT_ENV_VAR, raising=False)
    fallback = tmp_path / "data" / "one_week_backtest.xlsx"
    fallback.parent.mkdir(parents=True)
    fallback.write_bytes(b"x")
    assert resolve_input_path(tmp_path) == fallback


def test_resolve_input_path_ortam_degiskeni_yoksa_yedege_duser(
        tmp_path, monkeypatch):
    monkeypatch.setenv(INPUT_ENV_VAR, str(tmp_path / "olmayan.xlsx"))
    fallback = tmp_path / "data" / "one_week_backtest.xlsx"
    fallback.parent.mkdir(parents=True)
    fallback.write_bytes(b"x")
    assert resolve_input_path(tmp_path) == fallback


def test_resolve_input_path_hicbiri_yoksa_hata(tmp_path, monkeypatch):
    monkeypatch.delenv(INPUT_ENV_VAR, raising=False)
    with pytest.raises(InputContractError):
        resolve_input_path(tmp_path)


# ------------------------------------------------------------------- okuma

def test_referans_girdide_kimlik_eslemesi_ozdesliktir(data):
    frame, id_map, _notes = read_demand_table(REFERENCE_INPUT, data)
    assert list(frame.columns) == FORECAST_COLS
    assert len(frame) == 4046
    assert all(canonical == original
               for canonical, original in id_map.items())
    # Kanonik sıra src.forecast.assign_talep_ids ile birebir aynı olmalı.
    assert frame["Talep ID"].tolist() == [
        f"D{i:05d}" for i in range(1, len(frame) + 1)]


def test_referans_girdi_tahmin_ciktisiyla_ayni_degerleri_verir(data):
    frame, _id_map, _notes = read_demand_table(REFERENCE_INPUT, data)
    raw = pd.read_excel(REFERENCE_INPUT)
    assert int(frame["Tahmin Edilen Desi"].sum()) == int(
        raw["Tahmin Edilen Desi"].sum())
    assert set(frame["Talep Tamamlama Saati"]) == {"09:00", "17:00"}


def test_kanonik_olmayan_kimlikler_eslenir_ve_geri_donusur(tmp_path, data):
    frame = _demand_frame([
        ["REQ_a", "02.01.2030", "09:00:00", "İstanbul", "Yalova", 100],
        ["REQ_b", "01.01.2030", "17:00:00", "İstanbul", "Yalova", 200],
        ["REQ_c", "01.01.2030", "09:00:00", "İstanbul", "Yalova", 300],
    ])
    path = _write(frame, tmp_path / "girdi.xlsx")
    canonical, id_map, _notes = read_demand_table(path, data)

    # Sıra: tarih, cikis, varis, slot
    assert canonical["Talep ID"].tolist() == ["D00001", "D00002", "D00003"]
    assert id_map == {"D00001": "REQ_c", "D00002": "REQ_b",
                      "D00003": "REQ_a"}
    assert canonical["Tahmin Edilen Desi"].tolist() == [300, 200, 100]


def test_kolon_adlari_unicode_ve_bosluk_toleransli(tmp_path, data):
    frame = _demand_frame([
        ["D00001", "01.01.2030", "09:00", "İstanbul", "Yalova", 100],
    ])
    frame.columns = ["  talep id ", "TARİH", "Talep  Tamamlama Saati",
                     "Çıkış Transfer Merkezi", "Varış Transfer Merkezi",
                     "tahmin edilen desi"]
    path = _write(frame, tmp_path / "girdi.xlsx")
    canonical, _id_map, _notes = read_demand_table(path, data)
    assert list(canonical.columns) == FORECAST_COLS
    assert canonical["Tahmin Edilen Desi"].tolist() == [100]


def test_hat_matrisinde_olmayan_satir_atlanir(tmp_path, data):
    frame = _demand_frame([
        ["D00001", "01.01.2030", "09:00", "İstanbul", "Yalova", 100],
        ["D00002", "01.01.2030", "09:00", "Atlantis", "Yalova", 500],
        ["D00003", "01.01.2030", "09:00", "İstanbul", "İstanbul", 500],
    ])
    path = _write(frame, tmp_path / "girdi.xlsx")
    canonical, _id_map, notes = read_demand_table(path, data)
    assert len(canonical) == 1
    assert any("hat matrisinde olmayan" in note for note in notes)


def test_bos_tablo_hata_verir(tmp_path, data):
    path = _write(_demand_frame([]), tmp_path / "girdi.xlsx")
    with pytest.raises(InputContractError):
        read_demand_table(path, data)


def test_islenebilir_satir_yoksa_hata(tmp_path, data):
    frame = _demand_frame([
        ["D00001", "gecersiz", "09:00", "İstanbul", "Yalova", 100],
    ])
    path = _write(frame, tmp_path / "girdi.xlsx")
    with pytest.raises(InputContractError):
        read_demand_table(path, data)


# -------------------------------------------------------------------- ufuk

def test_ufuk_girdiden_turetilir_ve_kesintisizdir():
    frame = _demand_frame([
        ["D00001", "05.07.2027", "09:00", "İstanbul", "Yalova", 1],
        ["D00002", "02.07.2027", "09:00", "İstanbul", "Yalova", 1],
    ])
    days = horizon_days(frame)
    assert days[0] == dt.date(2027, 7, 2)
    assert days[-1] == dt.date(2027, 7, 5)
    assert len(days) == 4
    assert days == sorted(days)


def test_ufuk_tek_gun():
    frame = _demand_frame([
        ["D00001", "05.07.2027", "09:00", "İstanbul", "Yalova", 1],
    ])
    assert horizon_days(frame) == [dt.date(2027, 7, 5)]


def test_ufuk_kodda_gomulu_tarih_icermez():
    """Bölüm 7: aynı hafta farklı yıllara taşındığında sonuç kaymaz."""
    for year in (2019, 2026, 2031):
        frame = _demand_frame([
            ["D00001", f"29.06.{year}", "09:00", "İstanbul", "Yalova", 1],
            ["D00002", f"05.07.{year}", "17:00", "İstanbul", "Yalova", 1],
        ])
        days = horizon_days(frame)
        assert days[0] == dt.date(year, 6, 29)
        assert days[-1] == dt.date(year, 7, 5)
        assert len(days) == 7


# ------------------------------------------------------------------- kimlik

def test_restore_demand_ids_bolunmus_parca_sonekini_korur():
    plan = pd.DataFrame({"Talep ID": ["D00001", "D00002-3", "", None]})
    out = restore_demand_ids(plan, {"D00001": "AAA", "D00002": "BBB"})
    assert out["Talep ID"].tolist() == ["AAA", "BBB-3", "", ""]


def test_restore_demand_ids_bilinmeyen_kimligi_bozmaz():
    plan = pd.DataFrame({"Talep ID": ["D99999", "X-1"]})
    out = restore_demand_ids(plan, {"D00001": "AAA"})
    assert out["Talep ID"].tolist() == ["D99999", "X-1"]


def test_restore_demand_ids_bos_eslemede_degistirmez():
    plan = pd.DataFrame({"Talep ID": ["D00001"]})
    assert restore_demand_ids(plan, {})["Talep ID"].tolist() == ["D00001"]


# -------------------------------------------------------------------- yazım

def _plan_frame() -> pd.DataFrame:
    return pd.DataFrame([
        {"Araç ID": "V0001", "Araç Tipi": "Spot", "Araç türü": "Kamyon",
         "Çıkış Transfer Merkezi": "İstanbul",
         "Varış Transfer Merkezi": "Yalova",
         "Çıkış Tarihi": "01.01.2030", "Çıkış Saati": "09:00",
         "Varış Tarihi": "01.01.2030", "Varış Saati": "10:00",
         "Talep ID": "D00001", "Taşınan Desi": 100,
         "Yolculuk süresi": 56, "Varış elleçleme süresi": 1,
         "Çıkış Elleçleme süresi": 1, "SLA cezası": 0.0,
         "Toplam maliyet": 1234.5678},
        {"Araç ID": "V0002", "Araç Tipi": "Kiralık", "Araç türü": "Tır",
         "Çıkış Transfer Merkezi": "İstanbul",
         "Varış Transfer Merkezi": "Yalova",
         "Çıkış Tarihi": "01.01.2030", "Çıkış Saati": "09:00",
         "Varış Tarihi": "01.01.2030", "Varış Saati": "10:00",
         "Talep ID": "", "Taşınan Desi": 0,
         "Yolculuk süresi": 56, "Varış elleçleme süresi": 0,
         "Çıkış Elleçleme süresi": 0, "SLA cezası": 0.0,
         "Toplam maliyet": 900.0},
    ], columns=PLAN_COLS)


def test_write_final_plan_semayi_birebir_yazar(tmp_path, data):
    destination = tmp_path / "Tasima-plani.xlsx"
    write_final_plan(_plan_frame(), destination, data, {"D00001": "TALEP_7"})

    reloaded = pd.read_excel(destination)
    assert list(reloaded.columns) == PLAN_COLS
    assert len(reloaded) == 2
    assert str(reloaded["Talep ID"].iloc[0]) == "TALEP_7"
    assert pd.isna(reloaded["Talep ID"].iloc[1])
    assert reloaded["Toplam maliyet"].iloc[0] == pytest.approx(1234.5678)


def test_write_final_plan_tek_sayfa_yazar(tmp_path, data):
    from openpyxl import load_workbook
    destination = tmp_path / "Tasima-plani.xlsx"
    write_final_plan(_plan_frame(), destination, data, {})
    assert len(load_workbook(destination).sheetnames) == 1


def test_write_final_plan_gecersiz_plani_reddeder(tmp_path, data):
    broken = _plan_frame()
    broken.loc[0, "Araç türü"] = "Uçak"
    with pytest.raises(ValueError):
        write_final_plan(broken, tmp_path / "Tasima-plani.xlsx", data, {})
    assert not (tmp_path / "Tasima-plani.xlsx").exists()


def test_write_final_plan_hedef_klasoru_olusturur(tmp_path, data):
    destination = tmp_path / "yeni" / "klasor" / "Tasima-plani.xlsx"
    write_final_plan(_plan_frame(), destination, data, {})
    assert destination.is_file()


# ------------------------------------------------- round-trip toleransı

def test_roundtrip_excel_hassasiyet_farkini_tolere_eder():
    written = _plan_frame()
    written.loc[0, "Toplam maliyet"] = 3347.8859374999997
    reloaded = written.copy()
    # Excel bu değeri diskten 3347.8859375 olarak geri verir.
    reloaded.loc[0, "Toplam maliyet"] = 3347.8859375
    assert contract._roundtrip_differences(written, reloaded) == []


def test_roundtrip_gercek_sayisal_farki_yakalar():
    written = _plan_frame()
    reloaded = written.copy()
    reloaded.loc[0, "Toplam maliyet"] = 1234.5679
    differences = contract._roundtrip_differences(written, reloaded)
    assert len(differences) == 1
    assert "Toplam maliyet" in differences[0]


def test_roundtrip_metin_farkini_yakalar():
    written = _plan_frame()
    reloaded = written.copy()
    reloaded.loc[0, "Çıkış Saati"] = "09:01"
    differences = contract._roundtrip_differences(written, reloaded)
    assert len(differences) == 1
    assert "Çıkış Saati" in differences[0]


def test_roundtrip_bos_talep_id_nan_ile_esittir():
    written = _plan_frame()
    reloaded = written.copy()
    reloaded.loc[1, "Talep ID"] = float("nan")   # Excel boş hücreyi böyle verir
    assert contract._roundtrip_differences(written, reloaded) == []


def test_write_final_plan_hassasiyet_farkina_ragmen_yayinlar(tmp_path, data):
    """Bölüm 10: uyarı yüzünden çıktısız kalınmaz."""
    frame = _plan_frame()
    frame.loc[0, "Toplam maliyet"] = 3347.8859374999997
    destination = tmp_path / "Tasima-plani.xlsx"
    write_final_plan(frame, destination, data, {})
    assert destination.is_file()
    assert len(pd.read_excel(destination)) == 2


def test_write_final_plan_atomiktir_gecici_dosya_birakmaz(tmp_path, data):
    destination = tmp_path / "Tasima-plani.xlsx"
    write_final_plan(_plan_frame(), destination, data, {})
    leftovers = [p.name for p in tmp_path.iterdir() if p.name.startswith(".")]
    assert leftovers == []


# --------------------------------------------------- numpy skaler toleransı

def test_numpy_skalerleri_cozumlenir():
    import numpy as np
    # pandas hücreleri numpy skaleri döndürür; np.int64 Python int'in alt
    # sınıfı değildir, bu yüzden ayrıca kapsanır.
    assert contract._to_hhmm(np.int64(9)) == "09:00"
    assert contract._to_hhmm(np.float64(0.375)) == "09:00"
    assert contract._to_date(np.int64(46202)) == dt.date(2026, 6, 29)
    assert contract._to_desi(np.int64(7)) == 7


def test_bool_saat_veya_tarih_olarak_yorumlanmaz():
    assert contract._to_hhmm(True) is None
    assert contract._to_date(True) is None


def test_uzun_ufuk_uyari_basar_ama_calisir(capsys):
    frame = _demand_frame([
        ["D00001", "01.01.2030", "09:00", "İstanbul", "Yalova", 1],
        ["D00002", "01.03.2030", "09:00", "İstanbul", "Yalova", 1],
    ])
    days = horizon_days(frame)
    assert days[0] == dt.date(2030, 1, 1)
    assert days[-1] == dt.date(2030, 3, 1)
    assert "UYARI" in capsys.readouterr().out


# ------------------------------------------------------------ boş hücreler

@pytest.mark.parametrize("value", [None, float("nan"), pd.NaT, pd.NA, "", "  "])
def test_bos_hucreler_none_dondurur(value):
    # pd.NaT datetime'ın örneğidir; tip kontrollerinden önce elenmelidir.
    assert contract._to_date(value) is None
    assert contract._to_hhmm(value) is None
    assert contract._to_desi(value) is None


def test_bos_tarih_hucresi_olan_satir_atlanir(tmp_path, data):
    frame = _demand_frame([
        ["D00001", "01.01.2030", "09:00", "İstanbul", "Yalova", 100],
        ["D00002", None, "09:00", "İstanbul", "Yalova", 200],
    ])
    path = _write(frame, tmp_path / "girdi.xlsx")
    canonical, _id_map, notes = read_demand_table(path, data)
    assert len(canonical) == 1
    assert any("geçersiz tarih" in note for note in notes)
    assert canonical["Tarih"].tolist() == ["01.01.2030"]


def test_tireli_kimlikler_icin_uyari_notu(tmp_path, data):
    frame = _demand_frame([
        ["TALEP-1", "01.01.2030", "09:00", "İstanbul", "Yalova", 100],
        ["TALEP-2", "01.01.2030", "17:00", "İstanbul", "Yalova", 100],
    ])
    path = _write(frame, tmp_path / "girdi.xlsx")
    _canonical, id_map, notes = read_demand_table(path, data)
    assert any("tire var" in note for note in notes)
    assert sorted(id_map.values()) == ["TALEP-1", "TALEP-2"]
