import os
from datetime import date, datetime, time
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

import src.export as export_module
from src.export import require_plan_total, write_forecast_xlsx, write_plan_xlsx
from src.schemas import (FORECAST_COLS, PLAN_COLS, validate_forecast_grid,
                         validate_plan)


@pytest.fixture
def data():
    return SimpleNamespace(
        demand=pd.DataFrame([{"cikis": "A", "varis": "B"}]),
        tms=["A", "B"],
        lanes={("A", "B"): object()},
        vehicles={"Tır": object()},
    )


def _forecast_df():
    return pd.DataFrame([
        {"Talep ID": "D00001", "Tarih": "29.06.2026",
         "Talep Tamamlama Saati": "09:00",
         "Çıkış Transfer Merkezi": "A", "Varış Transfer Merkezi": "B",
         "Tahmin Edilen Desi": 100.0},
        {"Talep ID": "D00002", "Tarih": "29.06.2026",
         "Talep Tamamlama Saati": "17:00",
         "Çıkış Transfer Merkezi": "A", "Varış Transfer Merkezi": "B",
         "Tahmin Edilen Desi": 200.0},
    ], columns=FORECAST_COLS)


def _plan_df():
    return pd.DataFrame([
        {"Araç ID": "V0001", "Araç Tipi": "Spot", "Araç türü": "Tır",
         "Çıkış Transfer Merkezi": "A", "Varış Transfer Merkezi": "B",
         "Çıkış Tarihi": "29.06.2026", "Çıkış Saati": "10:00",
         "Varış Tarihi": "29.06.2026", "Varış Saati": "11:00",
         "Talep ID": "D00001", "Taşınan Desi": 100.0,
         "Yolculuk süresi": 60, "Varış elleçleme süresi": 10,
         "Çıkış Elleçleme süresi": 10, "SLA cezası": 0.0,
         "Toplam maliyet": 1000.0},
    ], columns=PLAN_COLS)


def test_require_plan_total_accepts_referee_total():
    frame = _plan_df()

    require_plan_total(frame, 1000.0)


def test_require_plan_total_accepts_difference_within_tolerance():
    frame = _plan_df()

    require_plan_total(frame, 999.995)


def test_require_plan_total_accepts_exact_decimal_cent_boundary():
    frame = _plan_df()
    frame.loc[0, "Toplam maliyet"] = 1.0

    require_plan_total(frame, 0.99)


def test_require_plan_total_rejects_just_beyond_tolerance():
    frame = _plan_df()
    frame.loc[0, "Toplam maliyet"] = 1.0

    with pytest.raises(ValueError, match="Toplam maliyet uzlaşmıyor"):
        require_plan_total(frame, 0.9899)


def test_require_plan_total_rejects_mismatch():
    frame = _plan_df()

    with pytest.raises(ValueError) as error:
        require_plan_total(frame, 999.0)

    assert str(error.value) == (
        "Toplam maliyet uzlaşmıyor: "
        "beyan 1000.00 TL, hakem 999.00 TL"
    )


@pytest.mark.parametrize(
    "invalid", [float("nan"), float("inf"), float("-inf"), "not-a-number"])
def test_require_plan_total_rejects_invalid_declared_cell(invalid):
    frame = pd.concat([_plan_df(), _plan_df()], ignore_index=True)
    frame["Toplam maliyet"] = frame["Toplam maliyet"].astype(object)
    frame.loc[1, "Toplam maliyet"] = invalid
    original = frame.copy(deep=True)

    with pytest.raises(ValueError, match="Toplam maliyet.*sonlu"):
        require_plan_total(frame, 1000.0)

    pd.testing.assert_frame_equal(frame, original)


@pytest.mark.parametrize(
    "invalid", [float("nan"), float("inf"), float("-inf"), None,
                "not-a-number"])
def test_require_plan_total_rejects_invalid_referee_total(invalid):
    frame = _plan_df()

    with pytest.raises(ValueError, match="Hakem toplam maliyeti.*sonlu"):
        require_plan_total(frame, invalid)


@pytest.mark.parametrize(
    "invalid", [float("nan"), float("inf"), float("-inf"), -0.01, None,
                "not-a-number"])
def test_require_plan_total_rejects_invalid_tolerance(invalid):
    frame = _plan_df()

    with pytest.raises(ValueError, match="Tolerans.*negatif olmayan"):
        require_plan_total(frame, 999.0, tolerance=invalid)


def test_write_forecast_xlsx_round_trips_exact_template(
        tmp_path, data, monkeypatch):
    path = tmp_path / "nested" / "forecast.xlsx"
    calls = []
    real_validator = export_module.validate_forecast_grid

    def tracking_validator(*args, **kwargs):
        calls.append(args[0])
        return real_validator(*args, **kwargs)

    monkeypatch.setattr(
        export_module, "validate_forecast_grid", tracking_validator)
    result = write_forecast_xlsx(
        _forecast_df(), path, data, date(2026, 6, 29), date(2026, 6, 29))

    loaded = pd.read_excel(path)
    assert result == path
    assert list(loaded.columns) == FORECAST_COLS
    assert validate_forecast_grid(
        loaded, data, date(2026, 6, 29), date(2026, 6, 29)) == []
    assert len(calls) == 2

    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet["C2"].value == time(9, 0)
    assert sheet["C3"].value == time(17, 0)
    assert sheet["C2"].number_format == "h:mm"
    assert sheet["C3"].number_format == "h:mm"
    assert sheet["F2"].number_format == "0.000"
    assert sheet["F3"].number_format == "0.000"


@pytest.mark.parametrize(
    ("morning", "evening"),
    [
        (time(9, 0), time(17, 0)),
        ("09:00:00", "17:00:00"),
        (datetime(2026, 6, 29, 9), datetime(2026, 6, 29, 17)),
        (pd.Timestamp(2026, 6, 29, 9), pd.Timestamp(2026, 6, 29, 17)),
    ],
    ids=["time", "zero-second-string", "datetime", "timestamp"],
)
def test_write_forecast_xlsx_formats_every_valid_slot_representation(
        tmp_path, data, morning, evening):
    path = tmp_path / "forecast.xlsx"
    frame = _forecast_df()
    frame["Talep Tamamlama Saati"] = pd.Series(
        [morning, evening], dtype=object)

    write_forecast_xlsx(
        frame, path, data, date(2026, 6, 29), date(2026, 6, 29))

    sheet = load_workbook(path).active
    assert [sheet["C2"].value, sheet["C3"].value] == [time(9), time(17)]
    assert [sheet["C2"].number_format, sheet["C3"].number_format] == [
        "h:mm", "h:mm"]
    assert [sheet["F2"].number_format, sheet["F3"].number_format] == [
        "0.000", "0.000"]


def test_format_forecast_workbook_rejects_noncanonical_slot_clearly(tmp_path):
    path = tmp_path / "forecast.xlsx"
    frame = _forecast_df()
    frame.loc[0, "Talep Tamamlama Saati"] = "09:00:01"
    frame.to_excel(path, index=False)

    with pytest.raises(ValueError, match="Geçersiz forecast saati.*09:00:01"):
        export_module._format_forecast_workbook(path)


def test_write_forecast_xlsx_rejects_invalid_without_file(tmp_path, data):
    path = tmp_path / "forecast.xlsx"
    incomplete = _forecast_df().iloc[[0]].copy()

    with pytest.raises(ValueError, match="Eksik forecast hücresi"):
        write_forecast_xlsx(
            incomplete, path, data, date(2026, 6, 29), date(2026, 6, 29))

    assert not path.exists()


def test_write_plan_xlsx_round_trips_exact_template(tmp_path, data):
    path = tmp_path / "nested" / "plan.xlsx"
    result = write_plan_xlsx(_plan_df(), path, data)

    loaded = pd.read_excel(path)
    assert result == path
    assert list(loaded.columns) == PLAN_COLS
    assert validate_plan(loaded, data) == []


def test_write_plan_xlsx_matches_template_desi_number_format(tmp_path, data):
    """The official plan template formats 'Taşınan Desi' as 0.000."""
    path = tmp_path / "plan.xlsx"
    frame = _plan_df()
    write_plan_xlsx(frame, path, data)

    sheet = load_workbook(path).active
    desi_column = chr(ord("A") + PLAN_COLS.index("Taşınan Desi"))
    assert sheet[f"{desi_column}1"].value == "Taşınan Desi"
    for row in range(2, len(frame) + 2):
        assert sheet[f"{desi_column}{row}"].number_format == "0.000"
    assert validate_plan(pd.read_excel(path), data) == []


def test_write_plan_xlsx_rejects_invalid_without_file(tmp_path, data):
    path = tmp_path / "plan.xlsx"
    invalid = _plan_df()
    invalid.loc[0, "Araç ID"] = "bad"

    with pytest.raises(ValueError, match="Araç ID formatı geçersiz"):
        write_plan_xlsx(invalid, path, data)

    assert not path.exists()


def _publication_case(tmp_path, *, plan_exists=True, forecast_exists=True):
    staging = tmp_path / "staging"
    official = tmp_path / "official"
    staging.mkdir()
    official.mkdir()
    staged_plan = staging / "plan.xlsx"
    staged_forecast = staging / "forecast.xlsx"
    plan_destination = official / "plan.xlsx"
    forecast_destination = official / "forecast.xlsx"
    staged_plan.write_bytes(b"new-plan")
    staged_forecast.write_bytes(b"new-forecast")
    if plan_exists:
        plan_destination.write_bytes(b"old-plan")
    if forecast_exists:
        forecast_destination.write_bytes(b"old-forecast")
    return SimpleNamespace(
        staged_plan=staged_plan,
        staged_forecast=staged_forecast,
        plan_destination=plan_destination,
        forecast_destination=forecast_destination,
    )


@pytest.mark.parametrize("failed_backup", ["plan", "forecast"])
def test_publish_backup_failure_never_starts_replacement(
        tmp_path, monkeypatch, failed_backup):
    case = _publication_case(tmp_path)
    failed_source = (
        case.plan_destination
        if failed_backup == "plan" else case.forecast_destination)
    real_copy2 = export_module.shutil.copy2
    replacement_calls = []

    def fail_selected_backup(source, destination):
        if source == failed_source:
            raise OSError(f"{failed_backup} backup failed")
        return real_copy2(source, destination)

    def forbidden_replace(*args):
        replacement_calls.append(args)
        raise AssertionError("replacement must not start after backup failure")

    monkeypatch.setattr(export_module.shutil, "copy2", fail_selected_backup)
    monkeypatch.setattr(export_module.os, "replace", forbidden_replace)

    with pytest.raises(OSError, match=f"{failed_backup} backup failed"):
        export_module.publish_workbooks(
            case.staged_plan,
            case.staged_forecast,
            case.plan_destination,
            case.forecast_destination,
        )

    assert replacement_calls == []
    assert case.plan_destination.read_bytes() == b"old-plan"
    assert case.forecast_destination.read_bytes() == b"old-forecast"
    assert case.staged_plan.read_bytes() == b"new-plan"
    assert case.staged_forecast.read_bytes() == b"new-forecast"


def test_publish_first_replacement_failure_restores_original_state(
        tmp_path, monkeypatch):
    case = _publication_case(tmp_path)
    real_replace = export_module.os.replace

    def fail_first_replacement(source, destination):
        if source == case.staged_plan:
            raise OSError("plan publication failed")
        return real_replace(source, destination)

    monkeypatch.setattr(
        export_module.os, "replace", fail_first_replacement)

    with pytest.raises(OSError, match="plan publication failed"):
        export_module.publish_workbooks(
            case.staged_plan,
            case.staged_forecast,
            case.plan_destination,
            case.forecast_destination,
        )

    assert case.plan_destination.read_bytes() == b"old-plan"
    assert case.forecast_destination.read_bytes() == b"old-forecast"
    assert case.staged_plan.read_bytes() == b"new-plan"
    assert case.staged_forecast.read_bytes() == b"new-forecast"


@pytest.mark.parametrize(
    ("plan_exists", "forecast_exists"),
    [(True, False), (False, True)],
    ids=["plan-existing", "forecast-existing"],
)
def test_publish_second_failure_restores_mixed_original_state(
        tmp_path, monkeypatch, plan_exists, forecast_exists):
    case = _publication_case(
        tmp_path,
        plan_exists=plan_exists,
        forecast_exists=forecast_exists,
    )
    real_replace = export_module.os.replace

    def fail_second_replacement(source, destination):
        if source == case.staged_forecast:
            raise OSError("forecast publication failed")
        return real_replace(source, destination)

    monkeypatch.setattr(
        export_module.os, "replace", fail_second_replacement)

    with pytest.raises(OSError, match="forecast publication failed"):
        export_module.publish_workbooks(
            case.staged_plan,
            case.staged_forecast,
            case.plan_destination,
            case.forecast_destination,
        )

    assert case.plan_destination.exists() is plan_exists
    assert case.forecast_destination.exists() is forecast_exists
    if plan_exists:
        assert case.plan_destination.read_bytes() == b"old-plan"
    if forecast_exists:
        assert case.forecast_destination.read_bytes() == b"old-forecast"


@pytest.mark.parametrize(
    ("plan_exists", "forecast_exists", "failed_actions"),
    [
        (True, True, ("plan restore",)),
        (True, True, ("forecast restore",)),
        (False, False, ("plan unlink",)),
        (False, False, ("forecast unlink",)),
        (True, False, ("plan restore", "forecast unlink")),
    ],
    ids=[
        "one-plan-restore",
        "one-forecast-restore",
        "one-plan-unlink",
        "one-forecast-unlink",
        "multiple-restore-and-unlink",
    ],
)
def test_publish_reports_and_chains_every_rollback_failure(
        tmp_path, monkeypatch, plan_exists, forecast_exists, failed_actions):
    case = _publication_case(
        tmp_path,
        plan_exists=plan_exists,
        forecast_exists=forecast_exists,
    )
    real_replace = export_module.os.replace
    real_unlink = Path.unlink

    def controlled_replace(source, destination):
        source = Path(source)
        destination = Path(destination)
        if source == case.staged_forecast:
            raise OSError("forecast publication failed")
        action = None
        if source.name == "0.backup" and destination == case.plan_destination:
            action = "plan restore"
        elif (source.name == "1.backup"
              and destination == case.forecast_destination):
            action = "forecast restore"
        if action in failed_actions:
            raise PermissionError(f"{action} failed")
        return real_replace(source, destination)

    def controlled_unlink(path, *args, **kwargs):
        action = None
        if path == case.plan_destination:
            action = "plan unlink"
        elif path == case.forecast_destination:
            action = "forecast unlink"
        if action in failed_actions:
            raise PermissionError(f"{action} failed")
        return real_unlink(path, *args, **kwargs)

    monkeypatch.setattr(export_module.os, "replace", controlled_replace)
    monkeypatch.setattr(Path, "unlink", controlled_unlink)

    with pytest.raises(RuntimeError, match="Workbook rollback başarısız") as caught:
        export_module.publish_workbooks(
            case.staged_plan,
            case.staged_forecast,
            case.plan_destination,
            case.forecast_destination,
        )

    message = str(caught.value)
    assert message.count("PermissionError(") == len(failed_actions)
    for action in failed_actions:
        destination = (
            case.plan_destination
            if action.startswith("plan") else case.forecast_destination)
        assert f"{destination}: PermissionError('{action} failed')" in message
    assert isinstance(caught.value.__cause__, OSError)
    assert str(caught.value.__cause__) == "forecast publication failed"


def test_publish_success_plan_then_forecast(tmp_path, monkeypatch):
    staging = tmp_path / "staging"
    official = tmp_path / "official"
    staging.mkdir()
    official.mkdir()
    staged_plan = staging / "plan.xlsx"
    staged_forecast = staging / "forecast.xlsx"
    plan_destination = official / "plan.xlsx"
    forecast_destination = official / "forecast.xlsx"
    staged_plan.write_bytes(b"new-plan")
    staged_forecast.write_bytes(b"new-forecast")
    plan_destination.write_bytes(b"old-plan")
    forecast_destination.write_bytes(b"old-forecast")
    real_replace = export_module.os.replace
    calls = []

    def tracking_replace(source, destination):
        calls.append((source, destination))
        real_replace(source, destination)

    monkeypatch.setattr(export_module.os, "replace", tracking_replace)

    published = export_module.publish_workbooks(
        staged_plan,
        staged_forecast,
        plan_destination,
        forecast_destination,
    )

    assert calls == [
        (staged_plan, plan_destination),
        (staged_forecast, forecast_destination),
    ]
    assert published == (plan_destination, forecast_destination)
    assert plan_destination.read_bytes() == b"new-plan"
    assert forecast_destination.read_bytes() == b"new-forecast"


def test_publish_second_failure_restores_both(tmp_path, monkeypatch):
    staging = tmp_path / "staging"
    official = tmp_path / "official"
    staging.mkdir()
    official.mkdir()
    staged_plan = staging / "plan.xlsx"
    staged_forecast = staging / "forecast.xlsx"
    plan_destination = official / "plan.xlsx"
    forecast_destination = official / "forecast.xlsx"
    staged_plan.write_bytes(b"new-plan")
    staged_forecast.write_bytes(b"new-forecast")
    plan_destination.write_bytes(b"old-plan")
    forecast_destination.write_bytes(b"old-forecast")
    real_replace = export_module.os.replace

    def fail_second_staged_replace(source, destination):
        if source == staged_forecast:
            raise OSError("forecast publication failed")
        real_replace(source, destination)

    monkeypatch.setattr(
        export_module.os, "replace", fail_second_staged_replace)

    with pytest.raises(OSError, match="forecast publication failed"):
        export_module.publish_workbooks(
            staged_plan,
            staged_forecast,
            plan_destination,
            forecast_destination,
        )

    assert plan_destination.read_bytes() == b"old-plan"
    assert forecast_destination.read_bytes() == b"old-forecast"


def test_publish_failure_removes_absent_destinations(tmp_path, monkeypatch):
    staging = tmp_path / "staging"
    official = tmp_path / "official"
    staging.mkdir()
    staged_plan = staging / "plan.xlsx"
    staged_forecast = staging / "forecast.xlsx"
    plan_destination = official / "plan.xlsx"
    forecast_destination = official / "forecast.xlsx"
    staged_plan.write_bytes(b"new-plan")
    staged_forecast.write_bytes(b"new-forecast")
    real_replace = export_module.os.replace

    def fail_second_staged_replace(source, destination):
        if source == staged_forecast:
            raise OSError("forecast publication failed")
        real_replace(source, destination)

    monkeypatch.setattr(
        export_module.os, "replace", fail_second_staged_replace)

    with pytest.raises(OSError, match="forecast publication failed"):
        export_module.publish_workbooks(
            staged_plan,
            staged_forecast,
            plan_destination,
            forecast_destination,
        )

    assert not plan_destination.exists()
    assert not forecast_destination.exists()


def _symlink_or_simulate(link, target, monkeypatch):
    try:
        link.symlink_to(target)
    except (NotImplementedError, OSError):
        link.unlink(missing_ok=True)
        link.write_bytes(target.read_bytes())
        real_is_symlink = Path.is_symlink

        def simulated_is_symlink(path):
            if path == link:
                return True
            return real_is_symlink(path)

        monkeypatch.setattr(Path, "is_symlink", simulated_is_symlink)


def test_publish_rejects_staged_source_symlink(tmp_path, monkeypatch):
    staging = tmp_path / "staging"
    official = tmp_path / "official"
    staging.mkdir()
    official.mkdir()
    source_target = tmp_path / "source-target.xlsx"
    source_target.write_bytes(b"new-plan")
    staged_plan = staging / "plan.xlsx"
    _symlink_or_simulate(staged_plan, source_target, monkeypatch)
    staged_forecast = staging / "forecast.xlsx"
    staged_forecast.write_bytes(b"new-forecast")
    plan_destination = official / "plan.xlsx"
    forecast_destination = official / "forecast.xlsx"
    plan_destination.write_bytes(b"old-plan")
    forecast_destination.write_bytes(b"old-forecast")

    with pytest.raises(ValueError, match="symlink"):
        export_module.publish_workbooks(
            staged_plan,
            staged_forecast,
            plan_destination,
            forecast_destination,
        )

    assert plan_destination.read_bytes() == b"old-plan"
    assert forecast_destination.read_bytes() == b"old-forecast"


def test_publish_rejects_existing_destination_symlink(tmp_path, monkeypatch):
    staging = tmp_path / "staging"
    official = tmp_path / "official"
    staging.mkdir()
    official.mkdir()
    staged_plan = staging / "plan.xlsx"
    staged_forecast = staging / "forecast.xlsx"
    staged_plan.write_bytes(b"new-plan")
    staged_forecast.write_bytes(b"new-forecast")
    destination_target = tmp_path / "destination-target.xlsx"
    destination_target.write_bytes(b"old-plan")
    plan_destination = official / "plan.xlsx"
    _symlink_or_simulate(plan_destination, destination_target, monkeypatch)
    forecast_destination = official / "forecast.xlsx"
    forecast_destination.write_bytes(b"old-forecast")

    def reject_backup(*args):
        raise AssertionError("symlink validation must precede backup")

    monkeypatch.setattr(export_module.shutil, "copy2", reject_backup)

    with pytest.raises(ValueError, match="symlink"):
        export_module.publish_workbooks(
            staged_plan,
            staged_forecast,
            plan_destination,
            forecast_destination,
        )

    assert destination_target.read_bytes() == b"old-plan"
    assert forecast_destination.read_bytes() == b"old-forecast"


def test_publish_rejects_known_hardlink_alias(tmp_path):
    staging = tmp_path / "staging"
    official = tmp_path / "official"
    staging.mkdir()
    official.mkdir()
    staged_plan = staging / "plan.xlsx"
    staged_forecast = staging / "forecast.xlsx"
    staged_plan.write_bytes(b"new-plan")
    staged_forecast.write_bytes(b"new-forecast")
    plan_destination = official / "plan.xlsx"
    forecast_destination = official / "forecast.xlsx"
    os.link(staged_plan, plan_destination)

    with pytest.raises(ValueError, match="alias"):
        export_module.publish_workbooks(
            staged_plan,
            staged_forecast,
            plan_destination,
            forecast_destination,
        )


@pytest.mark.parametrize("linked_path", ["source", "destination"])
def test_publish_rejects_multilink_regular_file(tmp_path, linked_path):
    staging = tmp_path / "staging"
    official = tmp_path / "official"
    staging.mkdir()
    official.mkdir()
    staged_plan = staging / "plan.xlsx"
    staged_forecast = staging / "forecast.xlsx"
    staged_plan.write_bytes(b"new-plan")
    staged_forecast.write_bytes(b"new-forecast")
    plan_destination = official / "plan.xlsx"
    forecast_destination = official / "forecast.xlsx"
    if linked_path == "source":
        linked = staged_plan
    else:
        plan_destination.write_bytes(b"old-plan")
        linked = plan_destination
    os.link(linked, tmp_path / f"outside-{linked_path}.xlsx")
    assert linked.stat().st_nlink > 1

    with pytest.raises(ValueError, match="multi-link"):
        export_module.publish_workbooks(
            staged_plan,
            staged_forecast,
            plan_destination,
            forecast_destination,
        )


def test_publish_rejects_resolved_path_collision(tmp_path):
    staging = tmp_path / "staging"
    staging.mkdir()
    (staging / "subdir").mkdir()
    staged_plan = staging / "plan.xlsx"
    staged_forecast = staging / "forecast.xlsx"
    staged_plan.write_bytes(b"new-plan")
    staged_forecast.write_bytes(b"new-forecast")
    plan_destination = staging / "subdir" / ".." / "plan.xlsx"
    forecast_destination = staging / "published-forecast.xlsx"

    with pytest.raises(ValueError, match="farklı olmalı"):
        export_module.publish_workbooks(
            staged_plan,
            staged_forecast,
            plan_destination,
            forecast_destination,
        )
